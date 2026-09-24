"""Documents in, markdown out — the file half of "give me clean text".

/v1/scrape handles a PDF at a URL; /v1/parse handles the file a customer
already has. PDFs are read page by page; Word documents by paragraph and
table; HTML through the same converter the scraper uses; plain text and
markdown pass through. Pages are the unit of metering: a PDF's real page
count, and for everything else one page per 3,000 characters, rounded up,
so a 40-page Word document is not billed as a single page.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass

CHARS_PER_PAGE = 3_000

_KINDS = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "text/plain": "text",
    "text/markdown": "text",
}
_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
}


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Parsed:
    markdown: str
    pages: int
    kind: str


def document_kind(filename: str, content_type: str) -> str | None:
    """What kind of document this is, or None if it is not one we parse.

    The QUESTION, separated from the ANSWERING. `/v1/parse` is handed a file
    and an unsupported one is a genuine error, so `kind_of` raises. The scrape
    path asks the same question about every page it fetches, where "not a
    document" is the ORDINARY answer — a JSON API, an RSS feed, an image.
    Calling the raising form there turned every JSON endpoint into a 500 (9 Sep
    2026: httpbin.org/headers, and with it every JSON API in a crawl).

    One mapping underneath both, because two copies of "which extensions and
    content types are documents" would drift the first time a type was added.
    """
    lower = (filename or "").lower()
    for ext, kind in _EXTENSIONS.items():
        if lower.endswith(ext):
            return kind
    return _KINDS.get((content_type or "").split(";")[0].strip().lower())


def kind_of(filename: str, content_type: str) -> str:
    """As above, for callers to whom "not a document" is an error."""
    kind = document_kind(filename, content_type)
    if kind is None:
        raise UnsupportedDocument("Supported files: PDF, DOCX, HTML, TXT and Markdown.")
    return kind


def parse_bytes(filename: str, content_type: str, data: bytes, *, max_pages: int = 500) -> Parsed:
    kind = kind_of(filename, content_type)
    if kind == "pdf":
        return _pdf(data, max_pages)
    if kind == "docx":
        markdown = _docx(data)
    elif kind == "html":
        markdown = _html(data)
    else:
        markdown = data.decode("utf-8", errors="replace").strip()
    return Parsed(markdown=markdown, pages=page_equivalents(markdown), kind=kind)


def page_equivalents(text: str) -> int:
    return max(1, math.ceil(len(text) / CHARS_PER_PAGE))


def _pdf(data: bytes, max_pages: int) -> Parsed:
    from pypdf import PdfReader  # BSD-3

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedDocument("This PDF is encrypted. Remove the password and try again.")
        count = len(reader.pages)
        if count > max_pages:
            raise UnsupportedDocument(f"This PDF has {count} pages; the limit is {max_pages}.")
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
    except UnsupportedDocument:
        raise
    except Exception as exc:  # noqa: BLE001 - a malformed file is the caller's problem, reported plainly
        raise UnsupportedDocument(f"Could not read this PDF: {exc}") from exc
    markdown = "\n\n".join(t for t in texts if t)
    return Parsed(markdown=markdown, pages=max(1, count), kind="pdf")


def _docx(data: bytes) -> str:
    from docx import Document  # python-docx, MIT

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedDocument(f"Could not read this Word document: {exc}") from exc

    out: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            out.append("#" * min(int(level), 6) + " " + text)
        elif "list" in style:
            out.append("- " + text)
        else:
            out.append(text)
    for table in document.tables:
        rows = [[cell.text.strip().replace("|", "\\|") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        out.append("| " + " | ".join(rows[0]) + " |")
        out.append("|" + "---|" * width)
        out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n\n".join(out).strip()


def _html(data: bytes) -> str:
    from selectolax.parser import HTMLParser

    from engine.core.extract.structured import convert_to_markdown

    tree = HTMLParser(data.decode("utf-8", errors="replace"))
    return convert_to_markdown(tree.body or tree.root).strip()
