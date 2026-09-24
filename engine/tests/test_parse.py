"""/v1/parse: documents to markdown, metered per page."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient

from engine.api import billing, deps
from engine.api.app import app
from engine.core.credits import credits_for
from engine.core.models import Cost
from engine.core.parse import UnsupportedDocument, page_equivalents, parse_bytes
from engine.storage.repositories import ApiKey


def _pdf(text: str) -> bytes:
    """A one-page PDF with a real xref table, so pypdf can read it."""
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 100]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        None,  # content stream, filled below
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode()
    objects[3] = b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


TINY_PDF = _pdf("Kettle manual")


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("Warranty", level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Part", "Years"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Element", "2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_pages_are_real_pages() -> None:
    parsed = parse_bytes("manual.pdf", "application/pdf", TINY_PDF)
    assert parsed.kind == "pdf" and parsed.pages == 1 and "Kettle manual" in parsed.markdown


def test_docx_becomes_headings_paragraphs_and_a_table() -> None:
    parsed = parse_bytes("warranty.docx", "", _docx(["Two years on the element."]))
    assert parsed.kind == "docx"
    assert "# Warranty" in parsed.markdown and "Two years on the element." in parsed.markdown
    assert "| Part | Years |" in parsed.markdown and "| Element | 2 |" in parsed.markdown
    assert parsed.pages == 1


def test_html_and_text_pass_through_and_pages_are_equivalents() -> None:
    html = parse_bytes(
        "page.html", "text/html", b"<html><body><h1>Hello</h1><p>World</p></body></html>"
    )
    assert html.kind == "html" and "Hello" in html.markdown and "World" in html.markdown
    text = parse_bytes("notes.txt", "text/plain", b"x" * 6_500)
    assert text.pages == 3 and page_equivalents("") == 1 and page_equivalents("a" * 3_000) == 1


def test_unknown_types_and_broken_pdfs_are_refused_plainly() -> None:
    with pytest.raises(UnsupportedDocument, match="Supported files"):
        parse_bytes("archive.zip", "application/zip", b"PK")
    with pytest.raises(UnsupportedDocument, match="Could not read this PDF"):
        parse_bytes("broken.pdf", "application/pdf", b"%PDF-1.4 nope")


@pytest.fixture
def charges(monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
    seen: list[dict[str, Any]] = []

    async def fake_charge(
        key: ApiKey, *, endpoint: str, url: str | None, cost: Cost, job_id: str | None = None
    ) -> int:
        seen.append({"endpoint": endpoint, "cost": cost, "credits": credits_for(cost)})
        return 1

    monkeypatch.setattr(billing, "charge", fake_charge)
    return seen


def _client(credits: int) -> TestClient:
    async def fake_key() -> ApiKey:
        return ApiKey(
            id="key_parse",
            label="parse",
            scopes=["scrape"],
            rate_limit_rpm=1000,
            allow_js_exec=False,
            webhook_secret=None,
            active=True,
            owner_ref="owner_1",
            credits_remaining=credits,
        )

    app.dependency_overrides[deps.require_api_key] = fake_key
    return TestClient(app)


@pytest.fixture(autouse=True)
def _clear() -> Iterator[None]:
    yield
    app.dependency_overrides.clear()


def test_the_route_parses_and_charges_per_page(charges: list[dict[str, Any]]) -> None:
    with _client(10) as client:
        r = client.post(
            "/v1/parse",
            files={"file": ("manual.pdf", TINY_PDF, "application/pdf")},
            headers={"Authorization": "Bearer k"},
        )
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert data["pages"] == 1 and data["kind"] == "pdf" and "Kettle manual" in data["markdown"]
    assert data["cost"] == {"pdf_pages": 1, "credits": 1}
    assert (
        charges[0]["endpoint"] == "parse"
        and charges[0]["cost"].pdf_pages == 1
        and charges[0]["credits"] == 1
    )


def test_the_route_refuses_without_credits_and_for_unsupported_files(
    charges: list[dict[str, Any]],
) -> None:
    with _client(0) as client:
        r = client.post(
            "/v1/parse",
            files={"file": ("m.pdf", TINY_PDF, "application/pdf")},
            headers={"Authorization": "Bearer k"},
        )
    assert r.status_code == 402
    with _client(10) as client:
        r = client.post(
            "/v1/parse",
            files={"file": ("a.zip", b"PK", "application/zip")},
            headers={"Authorization": "Bearer k"},
        )
    assert r.status_code == 400 and "Supported files" in r.json()["error"]["message"]
    assert charges == []
